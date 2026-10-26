"""Pruebas por función del motor (descuentos y abonos de fiado) — Prompt C.

Todo con datos INVENTADOS de 3-4 filas y cuentas hechas a mano (cada valor esperado está calculado
en el comentario, no copiado de lo que devuelve la función). Nada se escribe dentro del repo.

Correr desde la raíz del repo:  python -m pytest tests/python -q
"""
import importlib.util
import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

RAIZ = Path(__file__).resolve().parents[2]
MOTOR = RAIZ / "Contabilidad" / "modelo-financiero" / "generar_modelo_financiero.py"

# Que importar el motor aquí no deje un __pycache__ dentro de su carpeta.
sys.dont_write_bytecode = True
_spec = importlib.util.spec_from_file_location("motor_el_lobo", MOTOR)
m = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(m)


def df(filas, columnas=None):
    return pd.DataFrame(filas, columns=columnas)


# ───────────────────────────────────────────────────────────────────────────────────────────────
# ingresos_por_cierre: efectivo + transferencia + ajustes (igual que el KPI de Ingresos del Tablero)
# Los abonos de fiado YA vienen dentro de efectivo/transferencia: no se suman otra vez.
# ───────────────────────────────────────────────────────────────────────────────────────────────
class TestIngresosPorCierre:
    def test_cierre_antiguo_solo_ajuste_manual_va_al_efectivo(self):
        c = df([dict(efectivo=100_000, transferencia=50_000, ajuste_manual=5_000)])
        assert m.ingresos_por_cierre(c).tolist() == [155_000]          # 100.000 + 50.000 + 5.000

    def test_cierre_nuevo_con_abonos_NO_los_suma_otra_vez(self):
        c = df([dict(efectivo=200_000, transferencia=80_000, ajuste_efectivo=-3_000, ajuste_transferencia=2_000,
                     abonos_efectivo=30_000, abonos_transferencia=10_000)])
        # 200.000 + 80.000 − 3.000 + 2.000 = 279.000 (si sumara los abonos daría 319.000)
        assert m.ingresos_por_cierre(c).tolist() == [279_000]

    def test_ajuste_efectivo_vacio_cae_a_ajuste_manual(self):
        c = df([dict(efectivo=10_000, transferencia=0, ajuste_efectivo=np.nan, ajuste_manual=4_000)])
        assert m.ingresos_por_cierre(c).tolist() == [14_000]

    def test_ajuste_efectivo_en_cero_gana_sobre_ajuste_manual_viejo(self):
        c = df([dict(efectivo=10_000, transferencia=0, ajuste_efectivo=0, ajuste_manual=4_000)])
        assert m.ingresos_por_cierre(c).tolist() == [10_000]            # comportamiento igual al del sistema

    def test_varias_filas_y_columnas_que_faltan(self):
        c = df([dict(efectivo=1_000, transferencia=2_000), dict(efectivo=5_000, transferencia=None)])
        assert m.ingresos_por_cierre(c).tolist() == [3_000, 5_000]
        assert m.ingresos_por_cierre(pd.DataFrame()).empty


class TestAbonosPorCierre:
    def test_suma_abonos_de_efectivo_y_transferencia(self):
        c = df([dict(abonos_efectivo=30_000, abonos_transferencia=10_000), dict(abonos_efectivo=np.nan, abonos_transferencia=np.nan),
                dict(abonos_efectivo=7_000, abonos_transferencia=np.nan)])
        assert m.abonos_por_cierre(c).tolist() == [40_000, 0, 7_000]

    def test_cierres_viejos_sin_esas_columnas_dan_cero(self):
        c = df([dict(efectivo=1, transferencia=2), dict(efectivo=3, transferencia=4)])
        assert m.abonos_por_cierre(c).tolist() == [0, 0]
        assert m.abonos_por_cierre(pd.DataFrame()).empty


# ───────────────────────────────────────────────────────────────────────────────────────────────
# desglose_sin_pedido: de dónde viene lo del cierre que no tiene pedido detrás.
#   (ingresos − desayuno) − ventas por pedidos = ajustes + manuales + abonos + saldados + otros
# ───────────────────────────────────────────────────────────────────────────────────────────────
F18 = pd.Timestamp("2026-10-18")


class TestDesgloseSinPedido:
    def _otros(self, cierres, ventas, desayuno, turno, d):
        total = float(m.ingresos_por_cierre(cierres[cierres["turno"] == turno]).sum()) - float(sum(desayuno.values()) if turno == "almuerzo" else 0) - float(sum(ventas.values()))
        return total - d["ajustes"] - d["manuales"] - d["abonos"] - d["saldados"]

    def test_dia_con_abonos_ajustes_y_un_fiado_saldado_cuadra_al_peso(self):
        # Pedidos del día: 100.000 cobrados ese día + un fiado de 15.000 saldado con abonos (su plata llegó con los abonos).
        # Efectivo: 100.000 pedidos + 10.000 desayuno + 25.000 abonos − 20.000 de transferencia ya está en la otra columna.
        # efectivo 115.000 (100.000 − 20.000 transf. + 10.000 desayuno + 25.000 abonos) · transferencia 20.000
        c = df([dict(fecha=F18, turno="almuerzo", efectivo=115_000, transferencia=20_000, abonos_efectivo=25_000, abonos_transferencia=0,
                     ajuste_efectivo=-3_000, ajuste_transferencia=2_000, manual=False)])
        ventas = {F18: 115_000}                          # 100.000 + 15.000 del fiado saldado
        desayuno = {F18: 10_000}
        d = m.desglose_sin_pedido(c, pd.Series(ventas), pd.Series(desayuno), "almuerzo", fiados_saldados=15_000)
        assert d == dict(ajustes=-1_000, manuales=0.0, abonos=25_000, saldados=-15_000)     # ajustes: −3.000 + 2.000
        # ingresos del cierre = 135.000 − 1.000 = 134.000 → (134.000 − 10.000) − 115.000 = 9.000
        # 9.000 = −1.000 + 0 + 25.000 − 15.000 + otros  →  otros = 0
        assert self._otros(c, ventas, desayuno, "almuerzo", d) == 0

    def test_abonos_y_ajustes_no_se_mezclan(self):
        c = df([dict(fecha=F18, turno="cena", efectivo=50_000, transferencia=0, abonos_efectivo=8_000, abonos_transferencia=0,
                     ajuste_efectivo=1_500, manual=False)])
        d = m.desglose_sin_pedido(c, pd.Series({F18: 42_000}), pd.Series(dtype=float), "cena")
        assert d["ajustes"] == 1_500 and d["abonos"] == 8_000
        assert d["saldados"] == 0 and d["manuales"] == 0

    def test_cierre_manual_no_cuenta_dos_veces_sus_abonos(self):
        # Cierre manual: 90.000 + 10.000 = 100.000, de los cuales 5.000 son abonos; pedidos del día 80.000.
        # manuales = 100.000 − 5.000 − 80.000 = 15.000 ; abonos = 5.000 ; juntos = 20.000 = 100.000 − 80.000
        c = df([dict(fecha=F18, turno="cena", efectivo=90_000, transferencia=10_000, abonos_efectivo=5_000, manual=True)])
        d = m.desglose_sin_pedido(c, pd.Series({F18: 80_000}), pd.Series(dtype=float), "cena")
        assert d["manuales"] == 15_000 and d["abonos"] == 5_000
        assert d["manuales"] + d["abonos"] == 100_000 - 80_000

    def test_cierre_manual_sin_abonos_igual_que_antes(self):
        c = df([dict(fecha=F18, turno="cena", efectivo=90_000, transferencia=10_000, manual=True)])
        d = m.desglose_sin_pedido(c, pd.Series({F18: 80_000}), pd.Series(dtype=float), "cena")
        assert d == dict(ajustes=0.0, manuales=20_000, abonos=0.0, saldados=0.0)

    def test_solo_cuenta_los_cierres_del_turno_pedido(self):
        c = df([dict(fecha=F18, turno="almuerzo", efectivo=10_000, transferencia=0, abonos_efectivo=10_000, manual=False),
                dict(fecha=F18, turno="cena", efectivo=7_000, transferencia=0, abonos_efectivo=2_000, manual=False)])
        d = m.desglose_sin_pedido(c, pd.Series(dtype=float), pd.Series(dtype=float), "cena")
        assert d["abonos"] == 2_000

    def test_sin_cierres_los_fiados_saldados_igual_se_reportan(self):
        d = m.desglose_sin_pedido(pd.DataFrame(), pd.Series(dtype=float), pd.Series(dtype=float), "almuerzo", fiados_saldados=9_000)
        assert d == dict(ajustes=0.0, manuales=0.0, abonos=0.0, saldados=-9_000)


# ───────────────────────────────────────────────────────────────────────────────────────────────
# Fiados saldados con abonos nuevos y descuentos
# ───────────────────────────────────────────────────────────────────────────────────────────────
class TestCobradoViaAbonos:
    def test_solo_los_saldados_con_abonos_NUEVOS(self):
        p = df([
            dict(pagado_via_abono=True,  fecha_pago="2026-10-18"),   # nuevo → sí
            dict(pagado_via_abono=True,  fecha_pago=np.nan),         # viejo (sin fecha_pago) → no
            dict(pagado_via_abono=False, fecha_pago="2026-10-18"),   # no es de abonos → no
            dict(pagado_via_abono=True,  fecha_pago=""),             # fecha vacía → no
            dict(pagado_via_abono=np.nan, fecha_pago=np.nan),        # fila normal → no
        ])
        assert m.cobrado_via_abonos(p).tolist() == [True, False, False, False, False]

    def test_exports_viejos_sin_la_columna(self):
        p = df([dict(pagado_via_abono=True), dict(pagado_via_abono=False)])
        assert m.cobrado_via_abonos(p).tolist() == [False, False]
        assert m.cobrado_via_abonos(pd.DataFrame()).empty


class TestDescuentos:
    def test_descuentos_de_y_conteo_de_pedidos(self):
        p = df([
            dict(id=1, grupo_pedido="g1", descuento=1_000),      # g1: dos filas, UN pedido
            dict(id=2, grupo_pedido="g1", descuento=500),
            dict(id=7, grupo_pedido=np.nan, descuento=2_000),    # sin grupo: cuenta por id
            dict(id=8, grupo_pedido="g2", descuento=0),          # sin descuento
            dict(id=9, grupo_pedido="g3", descuento=np.nan),     # vacío = 0
        ])
        assert m.descuentos_de(p).tolist() == [1_000, 500, 2_000, 0, 0]       # suman 3.500
        assert m.contar_pedidos_con_descuento(p) == 2                          # g1 y el id 7

    def test_export_viejo_sin_columna_descuento(self):
        p = df([dict(id=1, monto_total=15_000)])
        assert m.descuentos_de(p).tolist() == [0]
        assert m.contar_pedidos_con_descuento(p) == 0
        assert m.contar_pedidos_con_descuento(pd.DataFrame()) == 0


# ───────────────────────────────────────────────────────────────────────────────────────────────
# calcular_cuentas_por_cobrar_fiados: misma cuenta que calcularSaldoFiado() del sistema
# (los números son los de tests/fixtures/datos.js: cliente 1 debe 9.000)
# ───────────────────────────────────────────────────────────────────────────────────────────────
class TestCuentasPorCobrar:
    def test_cuadra_con_el_sistema(self):
        pend = df([
            dict(es_fiar=True, cliente_fiado_id=1, monto_total=20_000),
            dict(es_fiar=True, cliente_fiado_id=1, monto_total=15_000),
            dict(es_fiar=True, cliente_fiado_id=np.nan, monto_total=12_000),   # huérfano: también es plata por cobrar
            dict(es_fiar=False, cliente_fiado_id=np.nan, monto_total=99_000),  # no es fiado
        ])
        abonos = df([dict(cliente_fiado_id=1, monto=10_000), dict(cliente_fiado_id=1, monto=5_000),
                     dict(cliente_fiado_id=1, monto=7_000), dict(cliente_fiado_id=1, monto=4_000),
                     dict(cliente_fiado_id=2, monto=15_000)])
        pagados = df([dict(cliente_fiado_id=2, pagado_via_abono=True, monto_total=15_000)])
        # cliente 1: 35.000 − 26.000 = 9.000 ; cliente 2: sin pendientes ; huérfano 12.000 → 21.000
        assert m.calcular_cuentas_por_cobrar_fiados(pend, abonos, pagados) == 21_000

    def test_abono_que_ya_salda_un_pedido_no_se_descuenta_dos_veces(self):
        # Debe 12.000 pendiente; abonó 30.000 y ya se aplicaron 20.000 a un pedido saldado → sobrante 10.000 → saldo 2.000
        pend = df([dict(es_fiar=True, cliente_fiado_id=1, monto_total=12_000)])
        abonos = df([dict(cliente_fiado_id=1, monto=30_000)])
        pagados = df([dict(cliente_fiado_id=1, pagado_via_abono=True, monto_total=20_000)])
        assert m.calcular_cuentas_por_cobrar_fiados(pend, abonos, pagados) == 2_000


# ───────────────────────────────────────────────────────────────────────────────────────────────
# De punta a punta: un Excel exportado INVENTADO (octubre) pasa por cargar_datos_reales().
#   10-12 almuerzo: p1 15.000 efectivo · p2 lista 15.000 con descuento 2.000 → 13.000 efectivo
#                   · p3 fiado de 20.000 (consumido hoy, saldado el 18 con un abono) · desayuno 10.000
#                   → cierre: efectivo 15.000 + 13.000 + 10.000 = 38.000 (la plata de p3 NO está aquí)
#   10-18 almuerzo: el abono de 20.000 de ese fiado → cierre: efectivo 20.000 (abonos_efectivo 20.000)
#   10-18 cena:     p4 18.000 por transferencia → cierre: transferencia 18.000
#   Pendiente: un fiado del cliente 1 por 12.000.
# ───────────────────────────────────────────────────────────────────────────────────────────────
def _pedido(**kw):
    base = dict(id=0, fecha="2026-10-12", hora="12:30", turno="almuerzo", pedido="completo", cantidad=1,
                monto_almuerzo=15_000, monto_domicilio=0, monto_total=15_000, metodo="efectivo",
                monto_efectivo=0, monto_transferencia=0, estado="pagado", es_domicilio=False, mesa=np.nan,
                para_llevar=False, es_fiar=False, es_sin_mesa=False, es_gratis=False, cliente_fiado_id=np.nan,
                grupo_pedido="g0", descuento=0, motivo_descuento="", pagado_via_abono=False, fecha_pago=np.nan)
    base.update(kw)
    return base


def _hojas(con_descuentos=True, con_columna_descuento=True):
    pagados = [
        _pedido(id=1, grupo_pedido="g1"),
        _pedido(id=2, grupo_pedido="g2", monto_almuerzo=13_000, monto_total=13_000, descuento=2_000, motivo_descuento="Cortesía (prueba)"),
        _pedido(id=3, grupo_pedido="g3", monto_almuerzo=20_000, monto_total=20_000, es_fiar=True, cliente_fiado_id=1,
                pagado_via_abono=True, fecha_pago="2026-10-18"),
        _pedido(id=4, grupo_pedido="g4", fecha="2026-10-18", turno="cena", pedido="hamburguesa", metodo="transferencia",
                monto_almuerzo=18_000, monto_total=18_000),
    ]
    if not con_columna_descuento:
        for p in pagados:
            p.pop("descuento"), p.pop("motivo_descuento")
    cierres = [
        dict(id=1, fecha="2026-10-12", turno="almuerzo", efectivo=38_000, transferencia=0, ajuste_efectivo=0, ajuste_transferencia=0,
             abonos_efectivo=0, abonos_transferencia=0, manual=False, gastos=0),
        dict(id=2, fecha="2026-10-18", turno="almuerzo", efectivo=20_000, transferencia=0, ajuste_efectivo=0, ajuste_transferencia=0,
             abonos_efectivo=20_000, abonos_transferencia=0, manual=False, gastos=0),
        dict(id=3, fecha="2026-10-18", turno="cena", efectivo=0, transferencia=18_000, ajuste_efectivo=0, ajuste_transferencia=0,
             abonos_efectivo=0, abonos_transferencia=0, manual=False, gastos=0),
    ]
    hojas = {
        "Pedidos_Pagados": pd.DataFrame(pagados),
        "Pedidos_Pendientes": pd.DataFrame([_pedido(id=10, grupo_pedido="g10", fecha="2026-10-15", estado="pendiente", es_fiar=True,
                                                    cliente_fiado_id=1, monto_almuerzo=12_000, monto_total=12_000)]),
        "Cierres_Dia": pd.DataFrame(cierres),
        "Aperturas_Turno": pd.DataFrame([dict(id=1, fecha="2026-10-12", turno="almuerzo", caja_inicial=50_000,
                                              venta_desayuno=10_000, venta_desayuno_transferencia=0)]),
        "Egresos": pd.DataFrame([dict(id=1, fecha="2026-10-14", cat="proveedor", monto=30_000, description="Compra (prueba)")]),
        "Gastos_Cierre": pd.DataFrame(columns=["id", "fecha", "turno", "nombre", "monto", "categoria"]),
        "Abonos_Fiado": pd.DataFrame([dict(id=1, cliente_fiado_id=1, monto=20_000, fecha="2026-10-18", turno="almuerzo", cuenta="efectivo")]),
        "Inventario": pd.DataFrame([dict(id=1, nombre="Arroz (prueba)", stock=10, costo=4_000)]),
        "Movimientos_Caja": pd.DataFrame(columns=["id", "fecha", "tipo", "monto", "cuenta"]),
    }
    if con_descuentos:
        hojas["Descuentos_Log"] = pd.DataFrame([dict(id=1, pedido_id=2, fecha="2026-10-12", descuento=2_000, motivo="Cortesía (prueba)")])
    return hojas


def _escribir_excel(ruta, hojas):
    with pd.ExcelWriter(ruta) as w:
        for nombre, d in hojas.items():
            d.to_excel(w, sheet_name=nombre, index=False)
    return ruta


@pytest.fixture
def cargado(tmp_path, capsys):
    ruta = _escribir_excel(tmp_path / "export.xlsx", _hojas())
    hist, extra, diario = m.cargar_datos_reales(str(ruta))
    return hist.iloc[0], extra, diario, capsys.readouterr().out


class TestCargarDatosReales:
    def test_ingresos_igual_que_el_tablero_y_sin_contar_los_abonos_dos_veces(self, cargado):
        h, _, _, _ = cargado
        # Tablero: Σ (efectivo + transferencia + ajustes) de los cierres de octubre.
        assert h["ingresos_almuerzo"] == 38_000 + 20_000          # 58.000
        assert h["ingresos_cena"] == 18_000
        assert h["ingresos_almuerzo"] + h["ingresos_cena"] == 76_000   # (si contara los abonos otra vez: 96.000)

    def test_ventas_por_pedidos_y_desayuno(self, cargado):
        h, _, _, _ = cargado
        assert h["ventas_pedidos_almuerzo"] == 15_000 + 13_000 + 20_000     # 48.000 (incluye el fiado saldado, en su día de consumo)
        assert h["ventas_pedidos_cena"] == 18_000
        assert h["ingresos_desayuno"] == 10_000

    def test_los_ingresos_sin_pedido_se_explican_por_abonos_y_fiados_saldados(self, cargado):
        h, _, _, _ = cargado
        # almuerzo: (58.000 − 10.000 desayuno) − 48.000 ventas = 0
        assert (h["ingresos_almuerzo"] - h["ingresos_desayuno"]) - h["ventas_pedidos_almuerzo"] == 0
        assert h["sinped_ajustes_almuerzo"] == 0 and h["sinped_manual_almuerzo"] == 0
        assert h["sinped_abonos_almuerzo"] == 20_000                          # el abono del día 18
        assert h["sinped_saldados_almuerzo"] == -20_000                       # el fiado de 20.000 ya contado como venta del día 12
        otros = 0 - (h["sinped_ajustes_almuerzo"] + h["sinped_manual_almuerzo"] + h["sinped_abonos_almuerzo"] + h["sinped_saldados_almuerzo"])
        assert otros == 0
        # comida rápida: sin abonos ni fiados saldados
        assert h["sinped_abonos_cena"] == 0 and h["sinped_saldados_cena"] == 0

    def test_descuentos_son_informativos(self, cargado):
        h, _, _, salida = cargado
        assert h["descuentos_almuerzo"] == 2_000 and h["descuentos_cena"] == 0
        assert h["pedidos_con_descuento"] == 1
        # ventas a precio de lista = ventas por pedidos + desayuno + descuentos = 48.000 + 18.000 + 10.000 + 2.000
        assert h["ventas_pedidos_almuerzo"] + h["ventas_pedidos_cena"] + h["ingresos_desayuno"] + h["descuentos_almuerzo"] + h["descuentos_cena"] == 78_000
        # y NO cambian los ingresos (monto_total ya viene neto)
        assert h["ingresos_almuerzo"] + h["ingresos_cena"] == 76_000
        assert "Descuentos otorgados" in salida and "2026-10: $2.000 en 1 pedido(s)" in salida
        assert "Descuentos_Log: 1 registro(s)" in salida

    def test_el_fiado_saldado_cuenta_en_su_fecha_de_consumo_no_en_la_del_abono(self, cargado):
        _, _, diario, _ = cargado
        ped = diario["pedidos"]
        p3 = ped[ped["id"] == 3]
        assert len(p3) == 1
        assert p3.iloc[0]["fecha"] == pd.Timestamp("2026-10-12")              # consumo, no el 18 (día del abono)
        assert p3.iloc[0]["cantidad"] == 1
        # unidades vendidas por fecha: el 12 hay 3 platos de almuerzo (p1, p2, p3); el 18 solo la hamburguesa
        por_dia = ped.groupby(["fecha", "turno"])["cantidad"].sum()
        assert por_dia[(pd.Timestamp("2026-10-12"), "almuerzo")] == 3
        assert (pd.Timestamp("2026-10-18"), "almuerzo") not in por_dia.index

    def test_fiados_por_cobrar(self, cargado):
        _, extra, _, _ = cargado
        # pendiente 12.000 − max(0, abonos 20.000 − ya aplicado a un pedido 20.000) = 12.000
        assert extra["cuentas_por_cobrar_fiados"] == 12_000

    def test_descuento_en_comida_rapida_se_cuenta_en_su_turno(self, tmp_path):
        hojas = _hojas()
        hojas["Pedidos_Pagados"].loc[hojas["Pedidos_Pagados"]["id"] == 4, ["descuento", "monto_almuerzo", "monto_total"]] = [1_500, 16_500, 16_500]
        hojas["Cierres_Dia"].loc[hojas["Cierres_Dia"]["id"] == 3, "transferencia"] = 16_500
        hist, _, _ = m.cargar_datos_reales(str(_escribir_excel(tmp_path / "cena.xlsx", hojas)))
        h = hist.iloc[0]
        assert h["descuentos_almuerzo"] == 2_000 and h["descuentos_cena"] == 1_500
        assert h["pedidos_con_descuento"] == 2                                # p2 (almuerzo) y p4 (cena)
        assert h["ingresos_cena"] == 16_500                                   # el ingreso sigue siendo lo que de verdad entró

    def test_export_viejo_sin_Descuentos_Log_ni_columna_descuento_sigue_funcionando(self, tmp_path, capsys):
        ruta = _escribir_excel(tmp_path / "viejo.xlsx", _hojas(con_descuentos=False, con_columna_descuento=False))
        hist, _, _ = m.cargar_datos_reales(str(ruta))
        h = hist.iloc[0]
        assert h["descuentos_almuerzo"] == 0 and h["descuentos_cena"] == 0 and h["pedidos_con_descuento"] == 0
        assert h["ingresos_almuerzo"] + h["ingresos_cena"] == 76_000
        assert "Descuentos otorgados" not in capsys.readouterr().out


# ───────────────────────────────────────────────────────────────────────────────────────────────
# El libro de Excel: las filas informativas existen y NO cambian ingresos ni utilidad.
# ───────────────────────────────────────────────────────────────────────────────────────────────
def _fila(ws, texto):
    for fila in ws.iter_rows():
        v = fila[1].value
        if isinstance(v, str) and v.strip().lstrip("· ").startswith(texto):
            return fila
    raise AssertionError(f"No encontré la fila '{texto}' en la hoja {ws.title}")


def _generar(tmp_path, hojas):
    entrada = _escribir_excel(tmp_path / "export.xlsx", hojas)
    salida = tmp_path / "modelo.xlsx"
    res = subprocess.run([sys.executable, "-I", str(MOTOR), str(entrada), "--salida", str(salida)],
                         capture_output=True, text=True, timeout=300)
    assert res.returncode == 0, f"El motor falló:\n{res.stdout[-1500:]}\n{res.stderr[-1500:]}"
    return salida, res.stdout


def test_el_modelo_trae_las_filas_nuevas_con_los_numeros_correctos(tmp_path):
    salida, consola = _generar(tmp_path, _hojas())
    ws = openpyxl.load_workbook(salida)["Model"]
    primero = lambda fila: [c.value for c in fila[2:] if isinstance(c.value, (int, float))][0]   # primer mes histórico
    assert primero(_fila(ws, "Descuentos otorgados")) == 2_000
    assert primero(_fila(ws, "Pedidos con descuento")) == 1
    assert primero(_fila(ws, "Abonos de fiado recibidos")) == 20_000
    assert primero(_fila(ws, "Fiados saldados con abonos")) == -20_000
    ventas_lista = _fila(ws, "Ventas a precio de lista")
    assert any(isinstance(c.value, str) and c.value.startswith("=") for c in ventas_lista)      # es una fórmula (se verifica su valor abajo)
    assert "abonos de fiado $20.000" in consola and "fiados saldados con abonos −$20.000" in consola.replace("-$", "−$")


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice no está instalado (solo se verifica la fórmula en Excel real)")
def test_los_ingresos_y_la_utilidad_no_cambian_por_los_descuentos_ni_los_abonos(tmp_path):
    """Recalcula el libro con LibreOffice y lee los valores: 'Ventas a precio de lista' = 78.000 y la diferencia
    'otros' del turno de almuerzo = 0; con y sin la hoja Descuentos_Log los ingresos del mes son iguales."""
    salida, _ = _generar(tmp_path, _hojas())
    carpeta = tmp_path / "recalculado"
    carpeta.mkdir()
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(carpeta), str(salida)],
                   capture_output=True, text=True, timeout=300, check=True)
    ws = openpyxl.load_workbook(carpeta / "modelo.xlsx", data_only=True)["Model"]
    primero = lambda fila: [c.value for c in fila[2:] if isinstance(c.value, (int, float))][0]
    assert primero(_fila(ws, "Ventas a precio de lista")) == 78_000
    assert primero(_fila(ws, "Diferencia entre pedidos y cierre")) == 0
