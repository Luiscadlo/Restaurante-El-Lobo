"""Pruebas del motor financiero (Contabilidad/modelo-financiero/generar_modelo_financiero.py).

Primera red de seguridad: el motor corre en modo --demo sin errores y genera el libro
de Excel esperado. Todo se escribe en una carpeta temporal: NADA dentro del repo.
(Las pruebas por función del motor se agregan junto con el Prompt C.)

Correr desde la raíz del repo:  python -m pytest tests/python -q
"""
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

RAIZ = Path(__file__).resolve().parents[2]
MOTOR = RAIZ / "Contabilidad" / "modelo-financiero" / "generar_modelo_financiero.py"

HOJAS_ESPERADAS = [
    "Cover", "Outputs", "A_Resultado", "B_Ingresos", "C_Egresos", "E_Proyeccion",
    "F_Familia", "G_Extras", "Inputs", "Model", "Datos_Graficos",
]


def _contenido_carpeta(carpeta: Path):
    return sorted(p.name for p in carpeta.iterdir())


@pytest.fixture(scope="module")
def corrida_demo(tmp_path_factory):
    """Corre el motor una sola vez en modo demo y devuelve (resultado, ruta_xlsx, antes, despues)."""
    salida = tmp_path_factory.mktemp("motor") / "modelo_demo.xlsx"
    antes = _contenido_carpeta(MOTOR.parent)
    # -I: modo aislado de Python (no carga nada extraño del directorio actual).
    res = subprocess.run(
        [sys.executable, "-I", str(MOTOR), "--demo", "--salida", str(salida)],
        capture_output=True, text=True, timeout=300,
    )
    despues = _contenido_carpeta(MOTOR.parent)
    return res, salida, antes, despues


def test_el_motor_termina_sin_error(corrida_demo):
    res, _, _, _ = corrida_demo
    assert res.returncode == 0, f"El motor falló:\n{res.stdout[-1500:]}\n{res.stderr[-1500:]}"
    assert "Traceback" not in res.stderr


def test_genera_el_libro_con_las_hojas_esperadas(corrida_demo):
    _, salida, _, _ = corrida_demo
    assert salida.exists() and salida.stat().st_size > 10_000
    libro = openpyxl.load_workbook(salida)
    assert libro.sheetnames == HOJAS_ESPERADAS


def test_no_escribe_nada_dentro_del_repo(corrida_demo):
    _, _, antes, despues = corrida_demo
    assert antes == despues, "El motor dejó archivos nuevos en la carpeta del motor"
    assert "__pycache__" not in despues


def test_las_hojas_no_estan_vacias(corrida_demo):
    _, salida, _, _ = corrida_demo
    libro = openpyxl.load_workbook(salida)
    for nombre in ("Outputs", "A_Resultado", "B_Ingresos", "C_Egresos", "Inputs", "Model"):
        hoja = libro[nombre]
        assert hoja.max_row > 5 and hoja.max_column > 2, f"La hoja {nombre} quedó casi vacía"
